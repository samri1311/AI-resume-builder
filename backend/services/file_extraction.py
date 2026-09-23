# backend/services/file_extraction.py
"""Resume file import: turns an uploaded PDF or Word (.docx) resume into
plain text, so it can fill the exact same "Your Background" box the
paste-text flow already uses - and from there feed the exact same
parse_background() prompt (see ai_engine.py) as any other pasted text,
completely unchanged. This is pure text extraction, no AI involved -
nothing here decides what the text means, it just gets it out of the file.

No OCR: a scanned/image-only PDF (a photo of a resume, or a PDF with no
real text layer) won't yield anything usable - that's reported back as a
clear error rather than silently returning nothing.
"""

import io

from docx import Document
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


def _extract_pdf_text(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages).strip()


def _extract_docx_text(file_bytes: bytes) -> str:
    doc = Document(io.BytesIO(file_bytes))

    parts = [p.text for p in doc.paragraphs if p.text.strip()]

    # Some resumes lay out contact info or skills in a table rather than
    # plain paragraphs - python-docx's .paragraphs alone would miss that
    # text entirely, so tables are walked too.
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text.strip())

    return "\n".join(parts).strip()


def extract_resume_text(filename: str, file_bytes: bytes) -> dict:
    """Returns {"success": True, "data": "<plain text>"} on success, or
    {"success": False, "error": "<message safe to show the person>"} - same
    success/data/error shape the rest of the app's parsing functions use.
    """
    if not filename:
        return {"success": False, "error": "No file was provided."}

    extension = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if extension not in SUPPORTED_EXTENSIONS:
        return {
            "success": False,
            "error": "Unsupported file type - please upload a PDF or Word (.docx) file.",
        }

    if not file_bytes:
        return {"success": False, "error": "That file appears to be empty."}

    try:
        if extension == ".pdf":
            text = _extract_pdf_text(file_bytes)
        else:
            text = _extract_docx_text(file_bytes)
    except Exception as e:
        print("Error extracting text from uploaded resume file:", e)
        return {
            "success": False,
            "error": "Couldn't read that file - it may be corrupted or password-protected.",
        }

    if not text:
        return {
            "success": False,
            "error": (
                "Couldn't find any text in that file - if it's a scanned or "
                "image-based PDF, try pasting the text instead."
            ),
        }

    return {"success": True, "data": text}
