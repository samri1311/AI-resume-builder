# backend/tests/test_file_extraction.py
"""Tests for backend.services.file_extraction.extract_resume_text.

PDF fixtures are built with reportlab (already a project dependency, and
independent of pypdf - the library under test) so these exercise a real
PDF file, not a mocked one. DOCX fixtures are built with python-docx
itself, since there's no independent DOCX-writer already in the project;
this still exercises the real .docx zip/XML format, just built and read
with the same library.
"""

import io

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from backend.services.file_extraction import extract_resume_text


def _build_pdf_bytes(lines):
    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=A4)
    y = 800
    for line in lines:
        c.drawString(50, y, line)
        y -= 20
    c.save()
    return buffer.getvalue()


def _build_docx_bytes(paragraphs, table_rows=None):
    buffer = io.BytesIO()
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table_rows:
        table = doc.add_table(rows=0, cols=len(table_rows[0]))
        for row_values in table_rows:
            row = table.add_row()
            for cell, value in zip(row.cells, row_values):
                cell.text = value
    doc.save(buffer)
    return buffer.getvalue()


def test_extracts_text_from_a_real_pdf():
    pdf_bytes = _build_pdf_bytes(["Jordan Rivera", "Backend Engineer at Acme Corp"])

    result = extract_resume_text("resume.pdf", pdf_bytes)

    assert result["success"] is True
    assert "Jordan Rivera" in result["data"]
    assert "Backend Engineer at Acme Corp" in result["data"]


def test_extracts_text_from_a_real_docx():
    docx_bytes = _build_docx_bytes(["Jordan Rivera", "Backend Engineer at Acme Corp"])

    result = extract_resume_text("resume.docx", docx_bytes)

    assert result["success"] is True
    assert "Jordan Rivera" in result["data"]
    assert "Backend Engineer at Acme Corp" in result["data"]


def test_extracts_table_text_from_docx():
    # Some resumes lay skills/contact info out in a table rather than plain
    # paragraphs - .paragraphs alone would miss this entirely.
    docx_bytes = _build_docx_bytes(
        paragraphs=["Jordan Rivera"],
        table_rows=[["Python", "FastAPI", "SQL"]],
    )

    result = extract_resume_text("resume.docx", docx_bytes)

    assert result["success"] is True
    assert "Python" in result["data"]
    assert "FastAPI" in result["data"]


def test_rejects_unsupported_file_extension():
    result = extract_resume_text("resume.txt", b"Jordan Rivera")

    assert result["success"] is False
    assert "Unsupported file type" in result["error"]


def test_rejects_filename_with_no_extension():
    result = extract_resume_text("resume", b"Jordan Rivera")

    assert result["success"] is False
    assert "Unsupported file type" in result["error"]


def test_rejects_missing_filename():
    result = extract_resume_text("", b"Jordan Rivera")

    assert result["success"] is False
    assert result["error"] == "No file was provided."


def test_rejects_empty_file_bytes():
    result = extract_resume_text("resume.pdf", b"")

    assert result["success"] is False
    assert "empty" in result["error"].lower()


def test_handles_corrupted_pdf_without_crashing():
    result = extract_resume_text("resume.pdf", b"this is not a real pdf file")

    assert result["success"] is False
    assert "couldn't read" in result["error"].lower()


def test_handles_corrupted_docx_without_crashing():
    result = extract_resume_text("resume.docx", b"this is not a real docx file")

    assert result["success"] is False
    assert "couldn't read" in result["error"].lower()


def test_pdf_with_no_extractable_text_reports_a_clear_error():
    # A blank PDF page (no text drawn at all) is the closest we can get to
    # a scanned/image-only PDF without an actual scanned fixture - both
    # produce zero extractable text.
    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=A4)
    c.save()  # no drawString calls - a valid PDF with a blank page

    result = extract_resume_text("resume.pdf", buffer.getvalue())

    assert result["success"] is False
    assert "scanned" in result["error"].lower()


def test_file_type_check_is_case_insensitive():
    pdf_bytes = _build_pdf_bytes(["Jordan Rivera"])

    result = extract_resume_text("Resume.PDF", pdf_bytes)

    assert result["success"] is True
